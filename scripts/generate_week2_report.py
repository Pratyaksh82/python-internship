import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page setup - Standard 1 inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_run_font(run, name='Calibri', size=11, color=None, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 'Calibri', 24, RGBColor(26, 54, 93), bold=True)
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 'Calibri', 12, RGBColor(74, 85, 104), italic=True)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 'Calibri', 15, RGBColor(31, 78, 121), bold=True)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 'Calibri', 12.5, RGBColor(43, 108, 176), bold=True)
    return p

def add_p(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, 'Calibri', 11, RGBColor(45, 55, 72))
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(bold_prefix + ' ')
    set_run_font(r1, 'Calibri', 11, RGBColor(26, 32, 44), bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 'Calibri', 11, RGBColor(45, 55, 72))
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shd = parse_xml(r'<w:shd {} w:fill="F7FAFC"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    borders = parse_xml(r'<w:tcBorders {}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/><w:left w:val="single" w:sz="24" w:space="0" w:color="3182CE"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/></w:tcBorders>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_text)
    set_run_font(r, 'Consolas', 9.5, RGBColor(45, 55, 72))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(title, body):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shd = parse_xml(r'<w:shd {} w:fill="EDF2F7"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd)
    borders = parse_xml(r'<w:tcBorders {}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + '\n')
    set_run_font(r1, 'Calibri', 10.5, RGBColor(43, 108, 176), bold=True)
    r2 = p.add_run(body)
    set_run_font(r2, 'Calibri', 10, RGBColor(74, 85, 104), italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Document Header
add_title('Debugging and Troubleshooting Python Applications')
add_subtitle('Week 2 Milestone Task: Systematic Bug Identification, pdb Tracing, Root-Cause Analysis, and Optimization\nJunior Python Developer Virtual Internship — Yuva Learn Certification\nPrepared by: Pratyaksh (@Pratyaksh82)')

# Metadata Table
meta_tbl = doc.add_table(rows=5, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Candidate Name", "Pratyaksh"),
    ("Program & Role", "Junior Python Developer Certification (Yuva Learn)"),
    ("Submission Milestone", "Week 2: Debugging and Troubleshooting Python Applications"),
    ("Source Code Repository", "https://github.com/Pratyaksh82/python-internship.git"),
    ("Execution Environment", "Python 3.14+ (pdb, logging, standard library)")
]
for i, (k, v) in enumerate(meta_data):
    c0 = meta_tbl.cell(i, 0)
    c1 = meta_tbl.cell(i, 1)
    c0.paragraphs[0].text = k
    set_run_font(c0.paragraphs[0].runs[0], 'Calibri', 10, RGBColor(26, 54, 93), bold=True)
    c1.paragraphs[0].text = v
    set_run_font(c1.paragraphs[0].runs[0], 'Calibri', 10, RGBColor(45, 55, 72))
    shd = parse_xml(r'<w:shd {} w:fill="F7FAFC"/>'.format(nsdecls('w')))
    c0._tc.get_or_add_tcPr().append(shd)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Section 1
add_h1('1. Executive Summary & Debugging Objectives')
add_p(
    'Software engineering in professional enterprise environments is rarely confined to writing greenfield code; '
    'over seventy percent of an engineer’s lifecycle is spent diagnosing runtime anomalies, refactoring legacy codebases, '
    'and isolating subtle logical regressions. The Week 2 milestone of the Junior Python Developer Virtual Internship '
    'focuses directly on mastering rigorous debugging methodologies, stack trace inspection, dynamic runtime analysis '
    'using the standard Python Debugger (pdb), and systematic code refactoring.'
)
add_p(
    'To rigorously fulfill the internship curriculum, an enterprise e-commerce order processing and analytics engine '
    'was engineered to simulate eight high-impact real-world bugs spanning file I/O descriptor leaks, mutable parameter defaults, '
    'off-by-one indexing, in-place sequence mutation during iteration, zero-division hazards, reference mutation in shallow copies, '
    'variable scope shadowing, and non-serializable object payloads. Each issue was systematically reproduced, diagnosed '
    'using pdb execution stepping, resolved through defensive software patterns, and validated.'
)

# Section 2
add_h1('2. Systematic Debugging Methodology & Tools')
add_p(
    'Modern software debugging relies on a repeatable, scientific troubleshooting loop rather than unstructured trial and error:'
)
add_bullet('1. Symptom Reproduction:', 'Constructing deterministic test cases and scripts (week2_debugging/debug_session.py) that consistently provoke the failure condition without external side effects.')
add_bullet('2. Hypothesis Formulation:', 'Inspecting the call stack, lexical scoping rules, object reference graphs, and memory mutations to form falsifiable hypotheses regarding the root cause.')
add_bullet('3. Interactive pdb Inspection:', 'Injecting breakpoints via breakpoint() or pdb.set_trace(), using inspection commands (n for next line, s for step into, c for continue, p for print expression, and locals() to inspect stack frames).')
add_bullet('4. Surgical Code Refactoring:', 'Implementing clean, idiomatic solutions that address the root cause while maintaining interface stability and preventing regression.')
add_bullet('5. Automated Verification:', 'Executing automated diagnostics across both buggy and corrected scripts to confirm that errors are eliminated.')

# Section 3
add_h1('3. Comprehensive Bug Audit & Resolution Matrix')

audit_tbl = doc.add_table(rows=9, cols=4)
audit_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Bug ID", "Fault Category", "Error / Symptom", "Applied Resolution"]
for j, h in enumerate(headers):
    c = audit_tbl.cell(0, j)
    c.paragraphs[0].text = h
    set_run_font(c.paragraphs[0].runs[0], 'Calibri', 10, RGBColor(255, 255, 255), bold=True)
    shd = parse_xml(r'<w:shd {} w:fill="2B6CB0"/>'.format(nsdecls('w')))
    c._tc.get_or_add_tcPr().append(shd)

audit_rows = [
    ("Bug 1", "File I/O & Resource Leak", "FileNotFoundError & unclosed file descriptor", "Context manager (with open) + os.path.exists guard"),
    ("Bug 2", "Mutable Default Argument", "Shared state across distinct order calls", "Default to None + deepcopy / fresh list allocation"),
    ("Bug 3", "Off-by-One Loop Index", "IndexError: list index out of range", "Replaced range(1, len+1) with direct iteration"),
    ("Bug 4", "Sequence Mutation", "Silent skip of consecutive cancelled items", "Replaced list.remove() with list comprehension"),
    ("Bug 5", "Coercion & ZeroDivision", "ZeroDivisionError & ValueError on '$19.99'", "Regex price parser + len == 0 guard returning 0.0"),
    ("Bug 6", "Shallow Copy Side-Effect", "Source order records mutated to archived=True", "Replaced list() shallow copy with copy.deepcopy()"),
    ("Bug 7", "Scoping & Shadowing", "UnboundLocalError: local variable 'tax' referenced", "Early exit validation + upfront variable initialization"),
    ("Bug 8", "JSON Serialization", "TypeError: datetime not JSON serializable", "Custom DateTimeEncoder inheriting json.JSONEncoder")
]

for i, (b_id, cat, sym, res) in enumerate(audit_rows, start=1):
    c0 = audit_tbl.cell(i, 0)
    c1 = audit_tbl.cell(i, 1)
    c2 = audit_tbl.cell(i, 2)
    c3 = audit_tbl.cell(i, 3)
    c0.paragraphs[0].text = b_id
    c1.paragraphs[0].text = cat
    c2.paragraphs[0].text = sym
    c3.paragraphs[0].text = res
    set_run_font(c0.paragraphs[0].runs[0], 'Calibri', 9.5, RGBColor(26, 54, 93), bold=True)
    set_run_font(c1.paragraphs[0].runs[0], 'Calibri', 9, RGBColor(45, 55, 72))
    set_run_font(c2.paragraphs[0].runs[0], 'Calibri', 9, RGBColor(197, 48, 48))
    set_run_font(c3.paragraphs[0].runs[0], 'Calibri', 9, RGBColor(40, 116, 166))
    if i % 2 == 1:
        shd = parse_xml(r'<w:shd {} w:fill="F7FAFC"/>'.format(nsdecls('w')))
        c0._tc.get_or_add_tcPr().append(shd)
        c1._tc.get_or_add_tcPr().append(shd)
        c2._tc.get_or_add_tcPr().append(shd)
        c3._tc.get_or_add_tcPr().append(shd)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Section 4
add_h1('4. Deep-Dive Code Analysis & Before-After Transformations')

add_h2('4.1 Bug 1: File I/O Descriptor Leaks and Missing Files')
add_p(
    'In the initial implementation, files were opened directly with f = open(filepath, "r") without error guards. '
    'If the file did not exist, an unhandled FileNotFoundError crashed the process. Even if the file existed, '
    'any exception raised during json.loads() bypassed the close() call, leaking the operating system file handle.'
)
add_code_block(
    "# BEFORE (Buggy):\n"
    "def load_raw_orders(filepath):\n"
    "    f = open(filepath, 'r')\n"
    "    content = f.read()\n"
    "    return json.loads(content)\n\n"
    "# AFTER (Refactored & Safe):\n"
    "def load_raw_orders(filepath: str) -> List[Dict[str, Any]]:\n"
    "    if not os.path.exists(filepath):\n"
    "        logger.warning('Target file %s does not exist. Initializing empty dataset.', filepath)\n"
    "        return []\n"
    "    try:\n"
    "        with open(filepath, 'r', encoding='utf-8') as f:\n"
    "            return json.load(f)\n"
    "    except (json.JSONDecodeError, OSError) as e:\n"
    "        logger.error('Failed to parse %s: %s', filepath, e)\n"
    "        return []"
)

add_h2('4.2 Bug 2: Mutable Default Arguments')
add_p(
    'In Python, default arguments are evaluated once when the function definition is executed, not each time '
    'the function is called. When a mutable object like a list ([]) is used as a default value, mutations '
    'persist across every call that relies on the default parameter.'
)
add_code_block(
    "# BEFORE (Buggy):\n"
    "def create_order(order_id, customer, items=[]):\n"
    "    items.append({'name': 'Processing Fee', 'price': '1.50'})\n"
    "    return {'order_id': order_id, 'items': items}\n\n"
    "# AFTER (Fixed):\n"
    "def create_order(order_id: int, customer: str, items: Optional[List[Dict[str, Any]]] = None):\n"
    "    order_items = [] if items is None else copy.deepcopy(items)\n"
    "    order_items.append({'name': 'Processing Fee', 'price': 1.50})\n"
    "    return {'order_id': order_id, 'customer': customer.strip(), 'items': order_items}"
)

add_h2('4.3 Bug 3 & 4: Sequence Iteration, Indexing, and In-Place Mutations')
add_p(
    'Attempting to delete elements from a list during iteration modifies the list length and shifts subsequent '
    'indices leftward. As a result, the loop skips items immediately following a deleted item. '
    'Furthermore, manual 1-based indexing causes IndexError on the last element.'
)
add_code_block(
    "# BEFORE (Buggy in-place removal skips consecutive items):\n"
    "for order in orders:\n"
    "    if order.get('status') == 'CANCELLED':\n"
    "        orders.remove(order)\n\n"
    "# AFTER (Idiomatic list comprehension):\n"
    "def filter_cancelled_orders(orders: List[Dict[str, Any]]) -> List[Dict[str, Any]]:\n"
    "    return [order for order in orders if order.get('status') != 'CANCELLED']"
)

add_h2('4.4 Bug 5 & 6: Data Type Sanitization & Reference Decoupling')
add_p(
    'Prices received from external APIs frequently contain currency symbols or whitespace (e.g., "$89.99"). '
    'Direct float conversion raises ValueError. Furthermore, calling list(orders) produces a shallow copy '
    'where the nested dictionaries are shared by reference, causing snapshot modifications to corrupt active records.'
)
add_code_block(
    "# Sanitizing string currency into clean float:\n"
    "def parse_price(price_val: Any) -> float:\n"
    "    if isinstance(price_val, (int, float)): return float(price_val)\n"
    "    if isinstance(price_val, str):\n"
    "        cleaned = re.sub(r'[^\\d.]', '', price_val.strip())\n"
    "        if cleaned: return float(cleaned)\n"
    "    return 0.0\n\n"
    "# Deep copy snapshot isolating historical state:\n"
    "def archive_orders_snapshot(orders: List[Dict[str, Any]]) -> List[Dict[str, Any]]:\n"
    "    snapshot = copy.deepcopy(orders)\n"
    "    for item in snapshot: item['archived'] = True\n"
    "    return snapshot"
)

add_h2('4.5 Bug 7 & 8: Scoping and JSON Datetime Serialization')
add_p(
    'Variables bound inside conditional branches produce UnboundLocalError when accessed outside or prior to assignment. '
    'Additionally, native Python datetime objects cannot be serialized by standard json.dump(). '
    'A custom JSONEncoder subclass solves this gracefully.'
)
add_code_block(
    "class DateTimeEncoder(json.JSONEncoder):\n"
    "    def default(self, obj: Any) -> Any:\n"
    "        if isinstance(obj, datetime):\n"
    "            return obj.isoformat()\n"
    "        return super().default(obj)"
)

# Section 5
add_h1('5. Diagnostic Execution Verification & Test Log')
add_p(
    'All diagnostic scenarios were verified using week2_debugging/debug_session.py. '
    'Below is the captured transcript demonstrating error reproduction on buggy code and zero-defect execution on corrected code.'
)

add_code_block(
    "======================================================================\n"
    " WEEK 2: PYTHON DEBUGGING AND TROUBLESHOOTING DEMONSTRATION\n"
    "======================================================================\n"
    "--- [Diagnosing Bug 1: Missing File & Resource Leak] ---\n"
    "-> Buggy Code Exception: FileNotFoundError: [Errno 2] No such file\n"
    "-> Corrected Code Handled Gracefully: Returns []\n\n"
    "--- [Diagnosing Bug 2: Mutable Default Argument] ---\n"
    "-> Buggy State: Order 1 items: 2, Order 2 items: 2 (LEAKED!)\n"
    "-> Corrected State: Order 1 count: 1, Order 2 count: 1 (ISOLATED!)\n\n"
    "--- [Diagnosing Bug 3: Off-by-One Loop Indexing] ---\n"
    "-> Buggy Code Exception: IndexError: list index out of range\n"
    "-> Corrected Code Output: Processed 2 orders successfully.\n\n"
    "--- [Diagnosing Bug 4: Mutating List During Iteration] ---\n"
    "-> Buggy Iteration Result: Length = 2 (Skipped second cancelled!)\n"
    "-> Corrected Filter Result: Length = 1 (Accurately filtered!)\n\n"
    "--- [Diagnosing Bug 5: ZeroDivisionError & Price Parsing] ---\n"
    "-> Buggy Empty Dataset Exception: ZeroDivisionError: division by zero\n"
    "-> Corrected Average on Empty Dataset: 0.0\n\n"
    "--- [Diagnosing Bug 6: Shallow Copy Mutation] ---\n"
    "-> Buggy Mutation: Source mutated to archived = True\n"
    "-> Corrected Mutation: Source preserved as archived = False\n\n"
    "--- [Diagnosing Bug 7: UnboundLocalError / Scope Confusion] ---\n"
    "-> Buggy Code Exception: UnboundLocalError: cannot access local variable 'tax'\n"
    "-> Corrected Calculation: Tax = $2.5\n\n"
    "--- [Diagnosing Bug 8: Datetime JSON Serialization] ---\n"
    "-> Buggy Code Exception: TypeError: Object of type datetime is not JSON serializable\n"
    "-> Corrected Export Result: Success = True\n"
    "======================================================================\n"
    " ALL 8 BUGS SUCCESSFULLY REPRODUCED, ISOLATED, AND RESOLVED!\n"
    "======================================================================"
)

# Section 6
add_h1('6. Interactive Debugging with Python pdb and IDE Tooling')
add_p(
    'Python provides a built-in interactive source code debugger via the pdb module. It supports setting conditional '
    'breakpoints, stepping through source code line-by-line, inspecting stack frames, evaluating arbitrary Python expressions, '
    'and conducting post-mortem analysis when unhandled exceptions occur.'
)
add_bullet('pdb Breakpoint Invocation:', 'Engineers can insert programmatic breakpoints into Python code using breakpoint() (Python 3.7+) or import pdb; pdb.set_trace(). When execution reaches this statement, control halts and yields an interactive (Pdb) prompt.')
add_bullet('Execution Stepping (n vs s):', 'The n (next) command continues execution until the next line in the current function is reached, executing any called functions without stepping into them. Conversely, s (step) steps execution into the target function call, allowing micro-level inspection of nested logic.')
add_bullet('Stack Frame Navigation (u vs d):', 'The u (up) command moves the current frame one level up in the call stack towards the caller, while d (down) moves one level down towards the callee, enabling inspection of parameter values passed across module boundaries.')
add_bullet('Post-Mortem Debugging (pdb.pm()):', 'When an unhandled exception terminates an interactive script, executing pdb.pm() immediately launches post-mortem analysis at the exact stack frame where the exception was raised, preserving local variable states at the moment of failure.')

# Section 7
add_h1('7. Code Optimization and Refactoring Principles')
add_p(
    'Beyond eliminating immediate runtime errors, the corrected application incorporates architectural refactoring:'
)
add_bullet('Type Annotations (PEP 484):', 'Full typing annotations (List, Dict, Optional, Any) improve code readability, enable static analysis with mypy, and prevent type errors before runtime.')
add_bullet('Structured Logging:', 'Replaced raw print statements with Python logging module, outputting formatted timestamps and severity levels (INFO, WARNING, ERROR).')
add_bullet('Defensive Deep Copying:', 'Guaranteed absolute data immutability by copying input lists and nested dictionaries before mutation.')
add_bullet('Comprehensive Input Sanitization:', 'Sanitizing customer strings with strip() and string currency values with regex before mathematical calculation.')

# Section 8
add_h1('8. Conclusion & Week 3 Roadmap')
add_p(
    'The Week 2 milestone successfully established advanced debugging capabilities required of professional Python developers. '
    'Through pdb inspection, root-cause analysis, and defensive refactoring, eight classic failure modes were eradicated. '
    'The resulting codebase serves as a benchmark for robust, production-grade Python engineering.'
)

output_dir = r"c:\Users\praty\Desktop\certification\python\docs"
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, "Python_Debugging_Troubleshooting_Week2_Report.docx")
doc.save(output_path)

total_words = 0
for p in doc.paragraphs:
    total_words += len(p.text.split())
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                total_words += len(p.text.split())

print(f"Report generated successfully at: {output_path}")
print(f"Total word count: {total_words} words")
